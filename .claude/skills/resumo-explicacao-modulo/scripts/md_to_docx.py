#!/usr/bin/env python
"""
Converte um Markdown gerado pela skill resumo-explicacao-modulo (padrao
PROMPT_PADRAO_RESUMO.md) em um .docx com o mesmo estilo visual dos documentos
ja existentes no repositorio (Arial, headings coloridos, tabelas com grid,
callouts em caixa, blocos de codigo monoespacados).

Uso:
    python md_to_docx.py ENTRADA.md SAIDA.docx

Nao inventa conteudo: apenas re-renderiza o Markdown existente em .docx.
"""
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

CALLOUT_EMOJIS = ("🎯", "💡", "⚡", "🔑", "🔍", "🏗️", "🏗", "✅", "🔌")

COLOR_H1 = RGBColor(0x2E, 0x74, 0xB5)
COLOR_H2 = RGBColor(0x2E, 0x74, 0xB5)
COLOR_H3 = RGBColor(0x1F, 0x4D, 0x78)
COLOR_H4 = RGBColor(0x1F, 0x4D, 0x78)
COLOR_BODY = RGBColor(0x33, 0x33, 0x33)
COLOR_BYLINE = RGBColor(0x55, 0x55, 0x55)
COLOR_CALLOUT_BORDER = "2E75B6"
FILL_CALLOUT = "F5F5F5"
FILL_CODE = "F0F0F0"
FILL_TABLE_HEADER = "D6E4F0"
BORDER_TABLE_CELL = "CCCCCC"


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color=BORDER_TABLE_CELL, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:color"), color)
        el.set(qn("w:sz"), sz)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_table_borders(table, sz="4"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:color"), "auto")
        el.set(qn("w:sz"), sz)
        borders.append(el)
    tblPr.append(borders)


def set_paragraph_shading(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def set_paragraph_left_border(paragraph, color=COLOR_CALLOUT_BORDER, sz="4", space="10"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:color"), color)
    left.set(qn("w:sz"), sz)
    left.set(qn("w:space"), space)
    pbdr.append(left)
    pPr.append(pbdr)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\n]+?\*)")


def add_inline_runs(paragraph, text, base_color=COLOR_BODY, base_bold=False, base_size=11):
    """Parseia **bold**, `code` e *italic* dentro de um texto e adiciona runs."""
    if text == "":
        paragraph.add_run("")
        return
    parts = INLINE_RE.split(text)
    for part in parts:
        if part == "":
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(base_size - 1)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.bold = base_bold
        run.font.name = run.font.name or "Arial"
        run.font.size = run.font.size or Pt(base_size)
        if run.font.color.rgb is None:
            run.font.color.rgb = base_color


def add_body_paragraph(doc, text, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    add_inline_runs(p, text)
    return p


def add_heading(doc, level, text):
    style_map = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}
    color_map = {1: COLOR_H1, 2: COLOR_H2, 3: COLOR_H3, 4: COLOR_H4}
    size_map = {1: 16, 2: 14, 3: 12, 4: 11}
    p = doc.add_paragraph(style=style_map.get(level, "Heading 4"))
    p.paragraph_format.space_before = Pt(18 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(8 if level <= 2 else 4)
    add_inline_runs(
        p,
        text,
        base_color=color_map.get(level, COLOR_BODY),
        base_bold=True,
        base_size=size_map.get(level, 11),
    )
    return p


def add_list_item(doc, text, ordered=False):
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)
    return p


def add_code_lines(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.left_indent = Pt(18)
        set_paragraph_shading(p, FILL_CODE)
        run = p.add_run(line if line != "" else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_BODY


def add_callout(doc, lines):
    """lines: linhas ja sem o prefixo '> '. Pode conter um bloco ``` aninhado."""
    in_code = False
    code_buffer = []
    text_lines = []
    for raw in lines:
        if raw.strip() == "```":
            in_code = not in_code
            if not in_code and code_buffer:
                text_lines.append(("code", code_buffer))
                code_buffer = []
            continue
        if in_code:
            code_buffer.append(raw)
        else:
            text_lines.append(("text", raw))

    n_blocks = len(text_lines)
    for idx, (kind, content) in enumerate(text_lines):
        is_first = idx == 0
        is_last = idx == n_blocks - 1
        if kind == "text":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Pt(18)
            pf.right_indent = Pt(18)
            pf.space_before = Pt(6 if is_first else 0)
            pf.space_after = Pt(6 if is_last else 0)
            set_paragraph_shading(p, FILL_CALLOUT)
            set_paragraph_left_border(p)
            add_inline_runs(p, content)
        else:
            for j, code_line in enumerate(content):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Pt(18)
                pf.right_indent = Pt(18)
                pf.space_before = Pt(0)
                pf.space_after = Pt(6 if (is_last and j == len(content) - 1) else 0)
                set_paragraph_shading(p, FILL_CALLOUT)
                set_paragraph_left_border(p)
                run = p.add_run(code_line if code_line != "" else " ")
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_BODY


def add_table(doc, header, rows):
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    set_table_borders(table)
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(header):
        cell = hdr_cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline_runs(p, text, base_bold=True)
        set_cell_shading(cell, FILL_TABLE_HEADER)
        set_cell_borders(cell)
        set_cell_margins(cell)
    for row in rows:
        cells = table.add_row().cells
        for i in range(n_cols):
            text = row[i] if i < len(row) else ""
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_runs(p, text)
            set_cell_shading(cell, "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_separator(line):
    return re.match(r"^\|?[\s:|-]+\|?$", line.strip()) is not None


def setup_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLOR_BODY
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "Arial")


def render_header(doc, title, meta_lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = COLOR_H1

    for i, meta in enumerate(meta_lines):
        clean = meta.strip()
        is_bold_line = clean.startswith("**") and clean.endswith("**")
        if is_bold_line:
            clean = clean[2:-2]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(clean)
        run.font.name = "Arial"
        if is_bold_line:
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_H1
        else:
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = COLOR_BYLINE
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        raw_lines = f.read().splitlines()

    doc = docx.Document()
    setup_base_styles(doc)

    # ---- cabecalho ----
    idx = 0
    title = ""
    if raw_lines and raw_lines[0].startswith("# "):
        title = raw_lines[0][2:].strip()
        idx = 1

    meta_lines = []
    while idx < len(raw_lines) and raw_lines[idx].strip() != "---":
        if raw_lines[idx].strip() != "":
            meta_lines.append(raw_lines[idx])
        idx += 1
    if idx < len(raw_lines) and raw_lines[idx].strip() == "---":
        idx += 1

    render_header(doc, title, meta_lines)

    lines = raw_lines[idx:]
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "" or stripped == "---":
            i += 1
            continue

        if stripped.startswith("#### "):
            add_heading(doc, 4, stripped[5:].strip())
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, 3, stripped[4:].strip())
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, 2, stripped[3:].strip())
            i += 1
            continue
        if stripped.startswith("# "):
            add_heading(doc, 1, stripped[2:].strip())
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # pula o fechamento
            add_code_lines(doc, code_lines)
            continue

        if stripped.startswith(">"):
            block = []
            while i < n and lines[i].strip().startswith(">"):
                content = lines[i].strip()[1:]
                if content.startswith(" "):
                    content = content[1:]
                block.append(content)
                i += 1
            add_callout(doc, block)
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            header = split_table_row(table_lines[0])
            body_lines = table_lines[1:]
            if body_lines and is_table_separator(body_lines[0]):
                body_lines = body_lines[1:]
            rows = [split_table_row(r) for r in body_lines]
            add_table(doc, header, rows)
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            add_list_item(doc, m.group(1), ordered=False)
            i += 1
            continue
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            add_list_item(doc, m.group(1), ordered=True)
            i += 1
            continue

        add_body_paragraph(doc, stripped)
        i += 1

    doc.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python md_to_docx.py ENTRADA.md SAIDA.docx")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
    print(f"OK: {sys.argv[2]}")
